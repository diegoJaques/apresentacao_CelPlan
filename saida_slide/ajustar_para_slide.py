from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Inches
import os

SRC = '/tmp/claude-0/-home-user-apresentacao-CelPlan/6db495d0-abf9-5be2-90f2-5e8c04ce6aa7/images/1.webp'
OUT = '/home/user/apresentacao_CelPlan/saida_slide'
os.makedirs(OUT, exist_ok=True)

src = Image.open(SRC).convert('RGB')
W, H = src.size

def ajustar(ratio_w, ratio_h, altura=2160):
    """Escala o infografico inteiro (sem cortes) e completa o formato
    replicando as colunas/linhas das bordas -> emenda invisivel."""
    tw, th = altura * ratio_w // ratio_h, altura
    esc = min(tw / W, th / H)
    sw, sh = round(W * esc), round(H * esc)
    big = src.resize((sw, sh), Image.LANCZOS)
    big = big.filter(ImageFilter.UnsharpMask(radius=1.6, percent=55, threshold=3))

    cv = Image.new('RGB', (tw, th))
    ox, oy = (tw - sw) // 2, (th - sh) // 2
    cv.paste(big, (ox, oy))
    if ox > 0:                                   # bordas laterais
        cv.paste(big.crop((0, 0, 1, sh)).resize((ox, sh), Image.NEAREST), (0, oy))
        rx = tw - sw - ox
        cv.paste(big.crop((sw-1, 0, sw, sh)).resize((rx, sh), Image.NEAREST), (tw-rx, oy))
    if oy > 0:                                   # bordas superior/inferior
        cv.paste(cv.crop((0, oy, tw, oy+1)).resize((tw, oy), Image.NEAREST), (0, 0))
        by = th - sh - oy
        cv.paste(cv.crop((0, oy+sh-1, tw, oy+sh)).resize((tw, by), Image.NEAREST), (0, th-by))
    return cv, (sw, sh), (ox, oy)

for nome, rw, rh in (('16x9', 16, 9), ('4x3', 4, 3)):
    img, cont, off = ajustar(rw, rh)
    p = os.path.join(OUT, f'infografico_epi_{nome}.png')
    img.save(p, 'PNG', optimize=True)
    print(f'{nome}: final {img.size}  conteudo {cont}  margens {off}  '
          f'{os.path.getsize(p)/1e6:.2f} MB')

# ---- PPTX 16:9 com a imagem ocupando o slide inteiro (full bleed) ----
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])          # layout em branco
s.shapes.add_picture(os.path.join(OUT, 'infografico_epi_16x9.png'),
                     Emu(0), Emu(0), prs.slide_width, prs.slide_height)
pptx_path = os.path.join(OUT, 'infografico_epi_16x9.pptx')
prs.save(pptx_path)
print('pptx:', pptx_path, f'{os.path.getsize(pptx_path)/1e6:.2f} MB')
