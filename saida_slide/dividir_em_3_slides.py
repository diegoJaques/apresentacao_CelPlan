"""Divide o infografico de EPIs em 3 slides 16:9 de 2 passos cada.

Cada card e recortado pelo seu proprio bounding box no original e reampliado.
A escala e UNICA para os 6 cards (limitada pelo par mais largo, passos 3+4),
para que o corpo de texto tenha o mesmo tamanho nos tres slides.
"""
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Inches
import os

SRC = '/tmp/claude-0/-home-user-apresentacao-CelPlan/6db495d0-abf9-5be2-90f2-5e8c04ce6aa7/images/1.webp'
OUT = '/home/user/apresentacao_CelPlan/saida_slide'

CABECALHO = (0, 155)
RODAPE    = (946, 1010)
FILEIRA1  = (174, 551)      # passos 1-3
FILEIRA2  = (563, 935)      # passos 4-6
# cada card: (x0, x1, faixa vertical da sua fileira)
CARDS = {1: (20, 496, FILEIRA1), 2: (512, 944, FILEIRA1), 3: (952, 1516, FILEIRA1),
         4: (20, 550, FILEIRA2), 5: (560, 1048, FILEIRA2), 6: (1060, 1516, FILEIRA2)}
PARES = [(1, 2), (3, 4), (5, 6)]

TW, TH = 3840, 2160
FUNDO   = (254, 254, 254)
MARGEM  = 100               # margem lateral minima (par mais largo)
GUT_SRC = 16                # vao entre cards no original -> escalonado junto

src = Image.open(SRC).convert('RGB')
W, H = src.size
nitidez = lambda i: i.filter(ImageFilter.UnsharpMask(radius=1.6, percent=55, threshold=3))

def faixa(y0, y1):
    f = src.crop((0, y0, W, y1)).resize((TW, round((y1-y0) * TW/W)), Image.LANCZOS)
    return nitidez(f)

cab, rod = faixa(*CABECALHO), faixa(*RODAPE)
banda_y0, banda_h = cab.height, TH - cab.height - rod.height

# escala unica: limitada em largura pelo par mais largo e em altura pela banda
larg = lambda p: sum(CARDS[i][1] - CARDS[i][0] for i in p)
alt   = lambda p: max(CARDS[i][2][1] - CARDS[i][2][0] for i in p)
esc = min(min((TW - 2*MARGEM) / (larg(p) + GUT_SRC) for p in PARES),
          min((banda_h - 140)   /  alt(p)            for p in PARES))
gut = round(GUT_SRC * esc)
print(f'escala unica = {esc:.3f}x  (vao entre cards {gut} px)')

paginas = []
for n, par in enumerate(PARES, start=1):
    imgs = []
    for i in par:
        x0, x1, (y0, y1) = CARDS[i]
        c = src.crop((x0, y0, x1, y1))
        imgs.append(nitidez(c.resize((round((x1-x0)*esc), round((y1-y0)*esc)), Image.LANCZOS)))
    bloco_w = sum(i.width for i in imgs) + gut
    bloco_h = max(i.height for i in imgs)
    x = (TW - bloco_w) // 2
    y = banda_y0 + (banda_h - bloco_h) // 2

    cv = Image.new('RGB', (TW, TH), FUNDO)
    cv.paste(cab, (0, 0))
    for im in imgs:                       # cards alinhados pelo topo, como no original
        cv.paste(im, (x, y)); x += im.width + gut
    cv.paste(rod, (0, TH - rod.height))
    p = os.path.join(OUT, f'infografico_epi_16x9_2passos_slide{n}.png')
    cv.save(p, 'PNG', optimize=True); paginas.append(p)
    print(f'slide {n} (passos {par[0]}+{par[1]}): bloco {bloco_w}x{bloco_h}  '
          f'margem lateral {(TW-bloco_w)//2}  {os.path.getsize(p)/1e6:.2f} MB')

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for p in paginas:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(p, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
d = os.path.join(OUT, 'infografico_epi_16x9_3slides.pptx')
prs.save(d); print('pptx:', d, f'{os.path.getsize(d)/1e6:.2f} MB')
