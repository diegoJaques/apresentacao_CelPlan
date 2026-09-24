"""Divide o infografico de EPIs em 2 slides 16:9 de 3 passos cada.

As faixas (cabecalho, fileira de cards, rodape) sao recortadas do original e
reampliadas com Lanczos na mesma escala, sem cortar conteudo dos cards.
"""
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Inches
import os

SRC = '/tmp/claude-0/-home-user-apresentacao-CelPlan/6db495d0-abf9-5be2-90f2-5e8c04ce6aa7/images/1.webp'
OUT = '/home/user/apresentacao_CelPlan/saida_slide'

# Faixas medidas no original 1536x1024 (topo inclusivo, base exclusiva)
CABECALHO = (0, 155)
FILEIRAS  = [(174, 551), (563, 935)]     # passos 1-3 e passos 4-6
RODAPE    = (946, 1010)

TW, TH = 3840, 2160                      # 16:9 (13,333" x 7,5" @ 288 dpi)
FUNDO   = (254, 254, 254)

src = Image.open(SRC).convert('RGB')
W, H = src.size
esc = TW / W                             # 2.5x - as faixas ocupam a largura toda

def faixa(y0, y1):
    """Recorta a faixa em largura cheia e reamplia na escala do slide."""
    f = src.crop((0, y0, W, y1))
    f = f.resize((TW, round((y1 - y0) * esc)), Image.LANCZOS)
    return f.filter(ImageFilter.UnsharpMask(radius=1.6, percent=55, threshold=3))

cab = faixa(*CABECALHO)
rod = faixa(*RODAPE)
folga_topo = 300                         # mesma posicao de cards nos 2 slides

paginas = []
for i, (y0, y1) in enumerate(FILEIRAS, start=1):
    cards = faixa(y0, y1)
    cv = Image.new('RGB', (TW, TH), FUNDO)
    cv.paste(cab, (0, 0))                                  # cabecalho no topo
    cv.paste(cards, (0, cab.height + folga_topo))          # fileira de 3 passos
    cv.paste(rod, (0, TH - rod.height))                    # rodape sangrando
    p = os.path.join(OUT, f'infografico_epi_16x9_slide{i}.png')
    cv.save(p, 'PNG', optimize=True)
    paginas.append(p)
    print(f'slide {i}: cards {cards.size}  sobra inferior '
          f'{TH - rod.height - (cab.height + folga_topo + cards.height)} px  '
          f'{os.path.getsize(p)/1e6:.2f} MB')

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for p in paginas:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(p, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
d = os.path.join(OUT, 'infografico_epi_16x9_2slides.pptx')
prs.save(d)
print('pptx:', d, f'{os.path.getsize(d)/1e6:.2f} MB')
