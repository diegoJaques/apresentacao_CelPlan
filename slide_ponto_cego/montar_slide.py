# -*- coding: utf-8 -*-
"""Reconstroi o slide 'O ponto cego' com dados de 2025 e cards desambiguados.

Uma unica fonte de conteudo gera os dois artefatos:
  - slide_ponto_cego.pptx  -> texto nativo e editavel (fonte Lato)
  - slide_ponto_cego.html  -> renderizado em PNG 3840x2160 pelo Chromium

Geometria e paleta medidas no slide original (1200x675 = 90 px/pol):
margem 68 px, 4 colunas de 266,5 px, reguas verticais, caixa NR-6 e
regua do rodape em y 620. Tamanhos de fonte calibrados pela altura real
dos glifos (digito 37 px -> 51 px de corpo, etc.).
"""
import os, subprocess, html

OUT = os.path.dirname(os.path.abspath(__file__))

# ----------------------------------------------------------------- paleta
ROXO, TINTA, CINZA, CINZA_C = '6A2F74', '17141B', '4C4652', '847D8B'
REGUA, CAIXA, PAGINA = 'DDD8E0', 'F2EAF4', 'C5C1C8'
FONTE = 'Lato'

# --------------------------------------------------------------- conteudo
EYEBROW = 'O PONTO CEGO'
TITULO  = 'O problema quase nunca é o EPI. É a fiscalização do uso.'

# (periodo, numero, unidade, o que e, base de comparacao)
# a etiqueta de periodo deixa explicito que os dois primeiros cards vem do
# estudo do MTE (2025) e os dois ultimos da serie do SmartLab (2012-2024)
CARDS = [
    ('2025',      '806.011', '',  'acidentes de trabalho registrados',
                                  'recorde da série histórica'),
    ('2025',      '3.644',   '',  'mortes por acidente de trabalho',
                                  'uma morte notificada a cada 2h24'),
    ('2012–2024', '13,6',    '%', 'dos acidentes têm máquina ou equipamento como agente causador',
                                  'amputam 15× mais e matam 3× mais que os demais agentes'),
    ('2012–2024', '5,1',     '%', 'dos afastamentos são de alimentador de linha de produção',
                                  'a ocupação que mais afasta no país; motorista vem em 2º, com 3,7%'),
]

# (texto, negrito) - o negrito marca a citacao literal da norma e o remate
CAIXA_NR6 = [
    ('A NR-6 é direta: cabe ao empregador ', False),
    ('“exigir seu uso”', True),
    (' (item 6.6.1, alínea “b”) — sem ressalva para as horas em que ninguém está '
     'olhando. Na prática, a verificação é amostral: acontece nos minutos em que '
     'alguém do SESMT passa pelo local. ', False),
    ('O resto do turno é ponto cego.', True),
]

FONTES = ('Fontes: MTE, estudo técnico “Acidentes do trabalho no Brasil — 2016 a 2025” '
          '(abr/2026; base INSS e eSocial), para o total de 2025. Observatório de SST — '
          'SmartLab (MPT/OIT), série 2012–2024, para agente causador e ocupação. '
          'NR-6, item 6.6.1. Em 2025 o volume absoluto é recorde, enquanto a taxa de CAT '
          'caiu de 29,39 (2016) para 17,94 — efeito do crescimento do emprego formal.')

RODAPE, PAG = 'VERIFICAÇÃO AUTOMÁTICA DE EPI · CELLVISION (VAIDIO)', '02'
LOGO = '/home/user/apresentacao_CelPlan/logo/logo sem fundo1.png'

# ------------------------------------------------- geometria (px @ 90/pol)
M, PITCH, COL = 68, 266.5, 245
Y_EYEBROW, Y_TITULO = 56, 85
Y_TAG, Y_NUM, Y_LABEL, Y_SUB = 194, 212, 284, 330
Y_CAIXA, H_CAIXA = 400, 128
Y_FONTES, Y_REGUA, Y_RODAPE = 556, 620, 640
S_NUM, S_UNID, S_TIT, S_EYE = 51, 24, 42, 12.5
S_TAG, S_LABEL, S_SUB, S_CAIXA, S_FONTES, S_RODAPE = 10.5, 15, 13, 18, 12, 11.5

# =========================================================== 1) HTML/PNG
def gerar_html():
    e = html.escape
    cards = []
    for i, (periodo, num, unid, rotulo, sub) in enumerate(CARDS):
        x = M + i * PITCH
        regua = (f'<div class="vr" style="left:{x-16:.1f}px"></div>' if i else '')
        cards.append(f'''{regua}
    <div class="tag"   style="left:{x:.1f}px">{e(periodo)}</div>
    <div class="num"   style="left:{x:.1f}px">{e(num)}<span class="un">{e(unid)}</span></div>
    <div class="label" style="left:{x:.1f}px">{e(rotulo)}</div>
    <div class="sub"   style="left:{x:.1f}px">{e(sub)}</div>''')
    nr6 = ''.join(f'<b>{e(t)}</b>' if b else e(t) for t, b in CAIXA_NR6)
    return f'''<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8">
<title>O ponto cego</title><style>
  *{{margin:0;padding:0;box-sizing:border-box}}
  body{{width:1200px;height:675px;position:relative;background:#fff;
       font-family:'{FONTE}',sans-serif;-webkit-font-smoothing:antialiased}}
  .eyebrow{{position:absolute;left:{M}px;top:{Y_EYEBROW}px;font-size:{S_EYE}px;
           font-weight:700;letter-spacing:2px;color:#{ROXO}}}
  h1{{position:absolute;left:{M}px;top:{Y_TITULO}px;width:700px;font-size:{S_TIT}px;
     font-weight:700;line-height:45px;color:#{TINTA};letter-spacing:-.3px}}
  .vr{{position:absolute;top:{Y_TAG}px;width:1px;height:178px;background:#{REGUA}}}
  .tag{{position:absolute;top:{Y_TAG}px;font-size:{S_TAG}px;font-weight:700;
       letter-spacing:1.4px;color:#{CINZA_C}}}
  .num{{position:absolute;top:{Y_NUM}px;font-size:{S_NUM}px;font-weight:300;
       color:#{ROXO};line-height:1}}
  .un{{font-size:{S_UNID}px;font-weight:400;margin-left:2px}}
  .label{{position:absolute;top:{Y_LABEL}px;width:{COL}px;font-size:{S_LABEL}px;
         line-height:22px;color:#{CINZA}}}
  .sub{{position:absolute;top:{Y_SUB}px;width:{COL}px;font-size:{S_SUB}px;
       line-height:19px;color:#{CINZA_C}}}
  .caixa{{position:absolute;left:{M+3}px;top:{Y_CAIXA}px;width:{1132-M-3}px;
         height:{H_CAIXA}px;background:#{CAIXA};border-left:3px solid #{ROXO}}}
  .caixa p{{padding:26px 34px;font-size:{S_CAIXA}px;line-height:30px;color:#{CINZA}}}
  .caixa b{{color:#{TINTA};font-weight:700}}
  .fontes{{position:absolute;left:{M}px;top:{Y_FONTES}px;width:{1132-M}px;
          font-size:{S_FONTES}px;line-height:18px;color:#{CINZA_C}}}
  .regua{{position:absolute;left:{M}px;top:{Y_REGUA}px;width:{1132-M}px;height:1px;
         background:#{REGUA}}}
  .rodape{{position:absolute;left:{M}px;top:{Y_RODAPE}px;font-size:{S_RODAPE}px;
          letter-spacing:1.4px;color:#{CINZA_C}}}
  .logo{{position:absolute;right:{1200-1112}px;top:{Y_RODAPE-9}px;height:32px}}
  .pag{{position:absolute;left:1118px;top:{Y_RODAPE}px;font-size:{S_RODAPE}px;color:#{PAGINA}}}
</style></head><body>
  <div class="eyebrow">{e(EYEBROW)}</div>
  <h1>{e(TITULO)}</h1>
  {''.join(cards)}
  <div class="caixa"><p>{nr6}</p></div>
  <div class="fontes">{e(FONTES)}</div>
  <div class="regua"></div>
  <div class="rodape">{e(RODAPE)}</div>
  <img class="logo" src="file://{LOGO}">
  <div class="pag">{e(PAG)}</div>
</body></html>'''

def renderizar_png():
    # headless_shell: o chrome completo desconta a barra da janela e corta ~81 px
    ch = '/opt/pw-browsers/chromium_headless_shell-1194/chrome-linux/headless_shell'
    p = os.path.join(OUT, 'slide_ponto_cego.html')
    open(p, 'w').write(gerar_html())
    subprocess.run([ch, '--no-sandbox', '--disable-gpu',
                    '--hide-scrollbars', '--force-device-scale-factor=3.2',
                    '--window-size=1200,675',
                    f'--screenshot={os.path.join(OUT, "slide_ponto_cego.png")}',
                    f'file://{p}'], check=True, capture_output=True)
    return os.path.join(OUT, 'slide_ponto_cego.png')

# ============================================================== 2) PPTX
def gerar_pptx(png):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    cor = lambda h: RGBColor.from_string(h)
    pol = lambda px: Inches(px / 90)
    pt  = lambda px: Pt(px * 0.8)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    def txt(x, y, w, h, align=PP_ALIGN.LEFT):
        tf = sl.shapes.add_textbox(pol(x), pol(y), pol(w), pol(h)).text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.paragraphs[0].alignment = align
        return tf

    def run(p, t, tam, c, bold=False, esp=None):
        r = p.add_run(); r.text = t
        r.font.name, r.font.size, r.font.bold = FONTE, pt(tam), bold
        r.font.color.rgb = cor(c)
        if esp: r._r.get_or_add_rPr().set('spc', str(int(esp * 0.8 * 100)))

    def rect(x, y, w, h, c):
        s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, pol(x), pol(y), pol(w), pol(h))
        s.fill.solid(); s.fill.fore_color.rgb = cor(c)
        s.line.fill.background(); s.shadow.inherit = False

    tf = txt(M, Y_EYEBROW - 2, 400, 20)
    run(tf.paragraphs[0], EYEBROW, S_EYE, ROXO, bold=True, esp=2)

    tf = txt(M, Y_TITULO - 6, 700, 100)
    tf.paragraphs[0].line_spacing = 45 / S_TIT
    run(tf.paragraphs[0], TITULO, S_TIT, TINTA, bold=True)

    for i, (periodo, num, unid, rotulo, sub) in enumerate(CARDS):
        x = M + i * PITCH
        if i: rect(x - 16, Y_TAG, 0.8, 178, REGUA)
        tf = txt(x, Y_TAG - 2, COL, 18)
        run(tf.paragraphs[0], periodo, S_TAG, CINZA_C, bold=True, esp=1.4)
        tf = txt(x, Y_NUM - 8, COL, 60)
        run(tf.paragraphs[0], num, S_NUM, ROXO)
        if unid: run(tf.paragraphs[0], unid, S_UNID, ROXO)
        tf = txt(x, Y_LABEL - 4, COL, 44)
        tf.paragraphs[0].line_spacing = 22 / S_LABEL
        run(tf.paragraphs[0], rotulo, S_LABEL, CINZA)
        tf = txt(x, Y_SUB - 4, COL, 44)
        tf.paragraphs[0].line_spacing = 19 / S_SUB
        run(tf.paragraphs[0], sub, S_SUB, CINZA_C)

    rect(M + 3, Y_CAIXA, 1132 - M - 3, H_CAIXA, CAIXA)
    rect(M, Y_CAIXA, 3, H_CAIXA, ROXO)
    tf = txt(M + 37, Y_CAIXA + 22, 1132 - M - 75, H_CAIXA - 40)
    tf.paragraphs[0].line_spacing = 30 / S_CAIXA
    for t, b in CAIXA_NR6:
        run(tf.paragraphs[0], t, S_CAIXA, TINTA if b else CINZA, bold=b)

    tf = txt(M, Y_FONTES - 4, 1132 - M, 40)
    tf.paragraphs[0].line_spacing = 18 / S_FONTES
    run(tf.paragraphs[0], FONTES, S_FONTES, CINZA_C)

    rect(M, Y_REGUA, 1132 - M, 0.8, REGUA)
    tf = txt(M, Y_RODAPE - 3, 700, 20)
    run(tf.paragraphs[0], RODAPE, S_RODAPE, CINZA_C, esp=1.4)
    sl.shapes.add_picture(LOGO, pol(1112 - 48), pol(Y_RODAPE - 9), height=pol(32))
    tf = txt(1118, Y_RODAPE - 3, 30, 20)
    run(tf.paragraphs[0], PAG, S_RODAPE, PAGINA)

    d = os.path.join(OUT, 'slide_ponto_cego.pptx')
    prs.save(d)
    return d

if __name__ == '__main__':
    png = renderizar_png(); print('png :', png)
    print('pptx:', gerar_pptx(png))
